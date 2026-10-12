import io
from datetime import date, timedelta
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.forms import inlineformset_factory
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from customers.models import Customer
from insurance.models import InsuranceCompany, InsuranceRecord
from vehicles.models import Vehicle
from .admin import UploadColumnFormSet
from .models import UploadColumn, UploadReceipt, UploadTemplate
from .schema import validate_columns
from .services import convert


class BulkUploadTests(APITestCase):
    def setUp(self):
        self.upload_bytes = {}
        self.user = User.objects.create_user(username="bulk-agent", password="test-pass")
        self.client.force_authenticate(self.user)
        self.template = UploadTemplate.objects.get(name="Standard Customers")
        self.insurance = UploadTemplate.objects.get(name="Standard Insurance Records")
        self.company = InsuranceCompany.objects.create(name="Example Insurance")

    def file(self, headers, rows):
        # Preview/import bind to exact bytes. Regenerating XLSX can change ZIP
        # timestamps even with identical cells, causing timing-dependent failures.
        key = (tuple(headers), tuple(tuple(row) for row in rows))
        if key not in self.upload_bytes:
            workbook = Workbook()
            workbook.active.append(headers)
            for row in rows:
                workbook.active.append(row)
            stream = io.BytesIO()
            workbook.save(stream)
            self.upload_bytes[key] = stream.getvalue()
        stream = io.BytesIO(self.upload_bytes[key])
        stream.name = "clients.xlsx"
        return stream

    def submit(self, headers, rows, template=None, token=None):
        payload = {"template_id": (template or self.template).pk, "file": self.file(headers, rows)}
        if token:
            payload["preview_token"] = token
        return self.client.post(f'/api/bulk-upload/{"import" if token else "preview"}/', payload, format="multipart")

    def insurance_data(self, policy="POL-1", vehicle="MH12AB1234"):
        today = timezone.localdate()
        return ["Phone", "Vehicle Number", "Insurance Company", "Policy No", "Start Date", "Expiry Date", "Premium"], [
            "9876543210", vehicle, self.company.name, policy, today, today + timedelta(days=365), "1200.00"
        ]

    def test_authentication_and_read_only_configuration_api(self):
        self.client.force_authenticate(None)
        self.assertIn(self.client.get("/api/bulk-upload/templates/").status_code, [401, 403])
        self.client.force_authenticate(self.user)
        self.assertEqual(self.client.get("/api/bulk-upload/templates/").status_code, 200)
        self.assertEqual(self.client.post("/api/bulk-upload/templates/", {"name": "Changed"}).status_code, 405)
        self.assertEqual(self.client.patch(f"/api/bulk-upload/templates/{self.template.pk}/download/", {}).status_code, 405)

    def test_dynamic_header_alias_and_preview_leaves_database_unchanged(self):
        column = self.template.columns.get(field_name="phone")
        column.column_name = "Client Contact"
        column.aliases = "Telephone"
        column.save()
        response = self.submit([" NAME ", "telephone"], [["Alice", 9876543210]])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["created"], 1)
        self.assertEqual(Customer.objects.count(), 0)
        imported = self.submit([" NAME ", "telephone"], [["Alice", 9876543210]], token=response.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(Customer.objects.get().phone, "9876543210")
        self.assertEqual(UploadReceipt.objects.count(), 1)

    def test_download_follows_admin_columns_and_is_safe_text(self):
        column = self.template.columns.get(field_name="name")
        column.column_name = "=Customer"
        column.save()
        self.template.columns.get(field_name="address").delete()
        response = self.client.get(f"/api/bulk-upload/templates/{self.template.pk}/download/")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        headers = [c.value for c in workbook.active[1]]
        self.assertIn("=Customer", headers)
        self.assertNotIn("Address", headers)
        self.assertEqual(workbook.active["A1"].data_type, "s")
        self.assertEqual(workbook.active["B2"].number_format, "@")

    def test_required_header_and_empty_value_are_rejected(self):
        response = self.submit(["Name"], [["Alice"]])
        self.assertEqual(response.status_code, 400)
        response = self.submit(["Name", "Phone"], [["Alice", None]])
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["errors"][0]["row"], 2)
        self.assertEqual(Customer.objects.count(), 0)

    def test_default_value_optional_column_and_unknown_header(self):
        column = self.template.columns.get(field_name="address")
        column.default_value = "Pune"
        column.save()
        preview = self.submit(["Phone", "Unused"], [["0123456789", "ignore"]])
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertTrue(preview.data["warnings"])
        response = self.submit(["Phone", "Unused"], [["0123456789", "ignore"]], token=preview.data["preview_token"])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Customer.objects.get().phone, "0123456789")
        self.assertEqual(Customer.objects.get().address, "Pune")

    def test_duplicate_and_ambiguous_headers(self):
        for headers in (["Phone", " phone "], ["Phone", "Mobile Number"]):
            response = self.submit(headers, [["9876543210", "9876543210"]])
            self.assertEqual(response.status_code, 400)

    def test_formula_and_bad_email_report_actual_excel_rows(self):
        response = self.submit(["Phone", "Email"], [["9876543210", "ok@example.com"], [None, None],
                                                       ["9876543211", "bad-email"], ["=1234567890", ""]])
        self.assertEqual(response.status_code, 400)
        self.assertEqual([e["row"] for e in response.data["errors"]], [4, 5])
        self.assertEqual(Customer.objects.count(), 0)

    def test_existing_customer_is_skipped_without_overwriting(self):
        customer = Customer.objects.create(name="Alice", phone="9876543210", address="Original")
        response = self.submit(["Name", "Phone", "Address"], [["Alice", "9876543210", "Replacement"]])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["skipped"], 1)
        imported = self.submit(["Name", "Phone", "Address"], [["Alice", "9876543210", "Replacement"]], token=response.data["preview_token"])
        self.assertEqual(imported.status_code, 200)
        customer.refresh_from_db()
        self.assertEqual(customer.address, "Original")

    def test_import_requires_preview_and_rejects_changed_file_or_config(self):
        headers, rows = ["Phone"], [["9876543210"]]
        missing = self.client.post("/api/bulk-upload/import/", {"template_id": self.template.pk, "file": self.file(headers, rows)}, format="multipart")
        self.assertEqual(missing.status_code, 400)
        preview = self.submit(headers, rows)
        token = preview.data["preview_token"]
        self.assertEqual(self.submit(headers, [["9876543211"]], token=token).status_code, 400)
        column = self.template.columns.get(field_name="name")
        column.column_name = "Client Name"
        column.save()
        self.assertEqual(self.submit(headers, rows, token=token).status_code, 400)
        self.assertEqual(Customer.objects.count(), 0)

    def test_preview_is_user_bound_expires_and_cannot_be_replayed(self):
        headers, rows = ["Phone"], [["9876543210"]]
        preview = self.submit(headers, rows)
        token = preview.data["preview_token"]
        other = User.objects.create_user(username="other")
        self.client.force_authenticate(other)
        self.assertEqual(self.submit(headers, rows, token=token).status_code, 400)
        self.client.force_authenticate(self.user)
        with patch("django.core.signing.time.time", return_value=timezone.now().timestamp() + 1900):
            self.assertEqual(self.submit(headers, rows, token=token).status_code, 400)
        self.assertEqual(self.submit(headers, rows, token=token).status_code, 200)
        self.assertEqual(self.submit(headers, rows, token=token).status_code, 400)
        self.assertEqual(Customer.objects.count(), 1)

    def test_insurance_preview_and_import_reuse_existing_business_rules(self):
        headers, row = self.insurance_data()
        preview = self.submit(headers, [row], template=self.insurance)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(InsuranceRecord.objects.count(), 0)
        self.assertEqual(Customer.objects.count(), 0)
        self.assertEqual(Vehicle.objects.count(), 0)
        response = self.submit(headers, [row], template=self.insurance, token=preview.data["preview_token"])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(InsuranceRecord.objects.get().policy_number, "POL-1")
        duplicate = self.submit(headers, [row], template=self.insurance)
        self.assertEqual(duplicate.status_code, 400)

    def test_duplicate_policy_within_file_rolls_back_all_related_objects(self):
        headers, row = self.insurance_data()
        response = self.submit(headers, [row, row], template=self.insurance)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(response.data["errors"][0]["row"], 3)
        self.assertEqual(InsuranceRecord.objects.count(), 0)
        self.assertEqual(Customer.objects.count(), 0)
        self.assertEqual(Vehicle.objects.count(), 0)

    def test_changed_database_after_preview_cancels_whole_import(self):
        headers, row = self.insurance_data()
        _, row2 = self.insurance_data("POL-2", "MH12AB1235")
        preview = self.submit(headers, [row, row2], template=self.insurance)
        self.assertEqual(preview.status_code, 200)
        other_preview = self.submit(headers, [row2], template=self.insurance)
        self.assertEqual(self.submit(headers, [row2], template=self.insurance, token=other_preview.data["preview_token"]).status_code, 200)
        response = self.submit(headers, [row, row2], template=self.insurance, token=preview.data["preview_token"])
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(response.data["created"], 0)
        self.assertEqual(InsuranceRecord.objects.count(), 1)
        self.assertFalse(Vehicle.objects.filter(vehicle_number="MH12AB1234").exists())
        self.assertEqual(UploadReceipt.objects.count(), 1)

    def test_unknown_company_and_invalid_dates_are_rejected(self):
        headers, row = self.insurance_data()
        row[2] = "Unknown company"
        self.assertEqual(self.submit(headers, [row], template=self.insurance).status_code, 400)
        row[2] = self.company.name
        row[5] = "2020-01-01"
        self.assertEqual(self.submit(headers, [row], template=self.insurance).status_code, 400)

    def test_multiple_active_policies_for_one_vehicle_are_rejected(self):
        headers, row = self.insurance_data()
        _, second = self.insurance_data("POL-2")
        response = self.submit(headers, [row, second], template=self.insurance)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("Active policy already exists", str(response.data["errors"]))
        self.assertEqual(InsuranceRecord.objects.count(), 0)

    def test_conflicting_vehicle_owner_is_rejected_without_mutation(self):
        owner = Customer.objects.create(name="Owner", phone="9876543211")
        vehicle = Vehicle.objects.create(customer=owner, vehicle_number="MH12AB1234", vehicle_type="Car")
        headers, row = self.insurance_data()
        response = self.submit(headers, [row], template=self.insurance)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(Customer.objects.count(), 1)
        self.assertEqual(InsuranceRecord.objects.count(), 0)
        vehicle.refresh_from_db()
        self.assertEqual(vehicle.customer_id, owner.pk)

    def test_file_and_row_limits(self):
        with patch("bulk_upload.services.MAX_ROWS", 1):
            self.assertEqual(self.submit(["Phone"], [["9876543210"], ["9876543211"]]).status_code, 400)
        with patch("bulk_upload.services.MAX_FILE_BYTES", 10):
            self.assertEqual(self.submit(["Phone"], [["9876543210"]]).status_code, 400)
        from django.core.files.uploadedfile import SimpleUploadedFile
        response = self.client.post("/api/bulk-upload/preview/", {"template_id": self.template.pk,
                                   "file": SimpleUploadedFile("bad.xlsx", b"invalid")}, format="multipart")
        self.assertEqual(response.status_code, 400)

    def test_admin_rejects_required_removal_and_alias_collisions(self):
        columns = list(self.template.columns.all())
        with self.assertRaises(ValidationError):
            validate_columns("customers", [c for c in columns if c.field_name != "phone"])
        columns[0].aliases = "phone"
        with self.assertRaises(ValidationError):
            validate_columns("customers", columns)

    def test_admin_inline_formset_rejects_required_deletion(self):
        factory = inlineformset_factory(UploadTemplate, UploadColumn, formset=UploadColumnFormSet,
                                       fields=["column_name", "field_name", "aliases", "is_required", "default_value", "position", "is_active"], extra=0)
        columns = list(self.template.columns.all())
        data = {"columns-TOTAL_FORMS": str(len(columns)), "columns-INITIAL_FORMS": str(len(columns))}
        for index, column in enumerate(columns):
            prefix = f"columns-{index}-"
            for key in ["id", "column_name", "field_name", "aliases", "default_value", "position"]:
                data[prefix + key] = str(getattr(column, key))
            data[prefix + "template"] = str(self.template.pk)
            if column.is_required:
                data[prefix + "is_required"] = "on"
            data[prefix + "is_active"] = "on"
            if column.field_name == "phone":
                data[prefix + "DELETE"] = "on"
        formset = factory(data, instance=self.template, prefix="columns")
        self.assertFalse(formset.is_valid())
        self.assertIn("Required mappings", str(formset.non_form_errors()))

    def test_numeric_and_date_conversions(self):
        self.assertEqual(convert(9876543210.0, "text"), "9876543210")
        self.assertEqual(convert("31/12/2026", "date"), "2026-12-31")
        self.assertEqual(convert("02.10.2025", "date"), "2025-10-02")
        with self.assertRaises(ValueError):
            convert("31.02.2025", "date")
        self.assertEqual(convert(1e20, "text"), "100000000000000000000")
        with self.assertRaises(ValueError):
            convert("NaN", "decimal")

    def compact_data(self):
        return ["S. DATE", "E. DATE", "CUT.NAME", None, None, "VEHICAL NO", "CO NAME", "NET PR.", "TOTAL."], [
            "02.10.2025", "01.10.2026", "Alice", None, None, "mh12 ab1234", self.company.name, "1000.00", "1200.00"
        ]

    def test_compact_format_preview_import_and_duplicate_protection(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        headers, row = self.compact_data()
        preview = self.submit(headers, [row], template=template)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["warnings"], [])
        self.assertEqual(Customer.objects.count(), 0)
        self.assertEqual(Vehicle.objects.count(), 0)
        self.assertEqual(InsuranceRecord.objects.count(), 0)
        response = self.submit(headers, [row], template=template, token=preview.data["preview_token"])
        self.assertEqual(response.status_code, 200, response.data)
        record = InsuranceRecord.objects.get()
        self.assertEqual(record.customer.phone, "")
        self.assertEqual(record.customer.name, "Alice")
        self.assertEqual(record.vehicle.vehicle_number, "MH12AB1234")
        self.assertEqual(record.policy_start_date, date(2025, 10, 2))
        self.assertEqual(record.policy_expiry_date, date(2026, 10, 1))
        self.assertEqual(record.total_premium, 1200)
        self.assertEqual(record.discount, 200)
        self.assertEqual(record.net_premium, 1000)
        self.assertTrue(record.policy_number.startswith("BULK-"))
        self.assertEqual(self.submit(headers, [row], template=template).status_code, 400)
        self.assertEqual(InsuranceRecord.objects.count(), 1)

    def test_compact_customer_resolution_and_ownership(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        headers, row = self.compact_data()
        owner = Customer.objects.create(name="Alice", phone="9876543210")
        Vehicle.objects.create(customer=owner, vehicle_number="MH12AB1234", vehicle_type="Car")
        preview = self.submit(headers, [row], template=template)
        self.assertEqual(preview.status_code, 200, preview.data)
        response = self.submit(headers, [row], template=template, token=preview.data["preview_token"])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(InsuranceRecord.objects.get().customer_id, owner.pk)
        owner.refresh_from_db()
        self.assertEqual(owner.phone, "9876543210")
        row[2] = "Bob"
        self.assertEqual(self.submit(headers, [row], template=template).status_code, 400)
        self.assertEqual(Customer.objects.count(), 1)

    def test_shared_new_placeholder_creates_separate_customer_vehicles(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        headers, first = self.compact_data()
        first[2], first[5] = "AADARSH HP GAS AGENCY", "NEW(1916)"
        second = first.copy()
        second[2] = "SHRI VIGHNAHAR HP GAS"
        preview = self.submit(headers, [first, second], template=template)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["created"], 2)
        self.assertEqual(Vehicle.objects.count(), 0)
        self.assertEqual(len(preview.data["warnings"]), 2)
        imported = self.submit(headers, [first, second], template=template, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        records = list(InsuranceRecord.objects.select_related("vehicle", "customer"))
        self.assertEqual(len(records), 2)
        self.assertNotEqual(records[0].vehicle_id, records[1].vehicle_id)
        self.assertNotEqual(records[0].policy_number, records[1].policy_number)
        for record in records:
            self.assertTrue(record.vehicle.vehicle_number.startswith("NEW(1916)-"))
            self.assertEqual(record.vehicle.customer_id, record.customer_id)
            self.assertEqual(record.net_premium, 1000)
        numbers = set(Vehicle.objects.values_list("vehicle_number", flat=True))
        duplicate = self.submit(headers, [second, first], template=template)
        self.assertEqual(duplicate.status_code, 400, duplicate.data)
        self.assertEqual(set(Vehicle.objects.values_list("vehicle_number", flat=True)), numbers)

    def test_same_registration_for_different_customers_in_file_is_rejected(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        headers, first = self.compact_data()
        second = first.copy()
        second[2] = "Bob"
        response = self.submit(headers, [first, second], template=template)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(response.data["errors"][0]["row"], 3)
        self.assertIn("Excel row 2", str(response.data["errors"]))
        self.assertEqual(Vehicle.objects.count(), 0)

    def test_compact_invalid_rows_roll_back_whole_file(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        headers, row = self.compact_data()
        for name, net in [("", "1000"), ("Bob", "1201"), ("Bob", "-1"), ("Bob", "1000.001")]:
            invalid = row.copy()
            invalid[2], invalid[5], invalid[7] = name, "MH12AB1235", net
            response = self.submit(headers, [row, invalid], template=template)
            self.assertEqual(response.status_code, 400, response.data)
            self.assertEqual(response.data["errors"][0]["row"], 3)
            self.assertEqual(Customer.objects.count(), 0)
            self.assertEqual(Vehicle.objects.count(), 0)

    def test_compact_ambiguous_customer_names_require_review(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        Customer.objects.create(name="Alice", phone="9876543210")
        Customer.objects.create(name="Alice", phone="9876543211")
        headers, row = self.compact_data()
        response = self.submit(headers, [row], template=template)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("Multiple customers", str(response.data["errors"]))
        self.assertEqual(InsuranceRecord.objects.count(), 0)

    def test_compact_template_accepts_descriptive_headers(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        headers = ["Start Date", "End Date", "Customer Name", "Vehicle Number", "Insurance Company", "Net Premium", "Total Premium"]
        rows = [["02.10.2025", "01.10.2026", "Alice", "MH12AB1234", self.company.name, "1000", "1200"]]
        response = self.submit(headers, rows, template=template)
        self.assertEqual(response.status_code, 200, response.data)

    def test_standard_template_still_requires_phone_and_policy(self):
        headers, row = self.insurance_data()
        for index in (0, 3):
            response = self.submit(headers[:index] + headers[index + 1:], [row[:index] + row[index + 1:]], template=self.insurance)
            self.assertEqual(response.status_code, 400)

    def make_archived_compact_policy(self, policy="AUTO-OLD", vehicle_number="MH12AB1234"):
        customer = Customer.objects.create(name="Alice", phone="9876543210")
        vehicle = Vehicle.objects.create(customer=customer, vehicle_number=vehicle_number, vehicle_type="Car")
        record = InsuranceRecord.objects.create(customer=customer, vehicle=vehicle, insurance_company=self.company,
            policy_number=policy, policy_start_date=date(2025, 10, 2), policy_expiry_date=date(2026, 10, 1), total_premium=1200)
        return customer, vehicle, record

    def test_reimport_recovers_legacy_policy_and_links_without_restoring_entire_batch(self):
        from payments.models import Payment
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        customer, vehicle, record = self.make_archived_compact_policy()
        other_customer, other_vehicle, other_record = self.make_archived_compact_policy("AUTO-OTHER", "MH12AB1235")
        payment = Payment.objects.create(insurance_record=record, amount=100, payment_method="Cash")
        Customer.objects.all().delete()  # All records share one deletion batch.
        headers, row = self.compact_data()
        preview = self.submit(headers, [row], template=template)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["restored"], 1)
        self.assertEqual(preview.data["created"], 0)
        self.assertEqual(preview.data["sample"][0]["status"], "restored")
        self.assertEqual(Customer.objects.count(), 0)
        self.assertEqual(Vehicle.objects.count(), 0)
        self.assertEqual(InsuranceRecord.objects.count(), 0)
        imported = self.submit(headers, [row], template=template, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(imported.data["restored"], 1)
        restored = InsuranceRecord.objects.get()
        self.assertEqual(restored.pk, record.pk)
        self.assertEqual(restored.policy_number, "AUTO-OLD")
        self.assertEqual(restored.customer_id, customer.pk)
        self.assertEqual(restored.vehicle_id, vehicle.pk)
        self.assertEqual(restored.net_premium, 1000)
        self.assertEqual(UploadReceipt.objects.get().restored_count, 1)
        self.assertEqual(Customer.objects.count(), 1)
        self.assertEqual(Vehicle.objects.count(), 1)
        self.assertEqual(Payment.objects.count(), 0)
        payment.refresh_from_db()
        self.assertIsNotNone(payment.deleted_at)
        other_record.refresh_from_db()
        self.assertIsNotNone(other_record.deleted_at)

    def test_standard_reimport_restores_explicit_archived_policy_number(self):
        headers, row = self.insurance_data()
        preview = self.submit(headers, [row], template=self.insurance)
        self.assertEqual(self.submit(headers, [row], template=self.insurance, token=preview.data["preview_token"]).status_code, 200)
        record = InsuranceRecord.objects.get()
        record.delete()
        preview = self.submit(headers, [row], template=self.insurance)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["restored"], 1)
        imported = self.submit(headers, [row], template=self.insurance, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(InsuranceRecord.objects.get().pk, record.pk)

    def test_invalid_file_rolls_back_archived_link_and_policy_recovery(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        customer, vehicle, record = self.make_archived_compact_policy()
        customer.delete()
        headers, row = self.compact_data()
        invalid = row.copy()
        invalid[5], invalid[6] = "MH12AB1235", "Unknown insurer"
        preview = self.submit(headers, [row], template=template)
        response = self.submit(headers, [row, invalid], template=template)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(Customer.objects.count(), 0)
        self.assertEqual(Vehicle.objects.count(), 0)
        self.assertEqual(InsuranceRecord.objects.count(), 0)
        # State can change after a successful preview; import must roll back too.
        self.company.is_active = False
        self.company.save()
        imported = self.submit(headers, [row], template=template, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 400, imported.data)
        self.assertEqual(UploadReceipt.objects.count(), 0)
        record.refresh_from_db()
        self.assertIsNotNone(record.deleted_at)

    def test_archived_vehicle_owner_mismatch_uses_uploaded_customer(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        customer, vehicle, record = self.make_archived_compact_policy()
        customer.delete()
        headers, row = self.compact_data()
        row[2] = "Bob"
        response = self.submit(headers, [row], template=template)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["restored"], 1)
        self.assertIn("Alice", str(response.data["warnings"]))
        self.assertIn("Bob", str(response.data["warnings"]))
        self.assertEqual(Vehicle.objects.count(), 0)
        self.assertEqual(Customer.objects.count(), 0)
        vehicle.refresh_from_db()
        self.assertEqual(vehicle.customer_id, customer.pk)
        imported = self.submit(headers, [row], template=template, token=response.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        restored = InsuranceRecord.objects.get()
        self.assertEqual(restored.pk, record.pk)
        self.assertEqual(restored.customer.name, "Bob")
        self.assertEqual(restored.vehicle.customer_id, restored.customer_id)
        self.assertEqual(Customer.objects.count(), 1)
        customer.refresh_from_db()
        self.assertIsNotNone(customer.deleted_at)

    def test_reassigned_archived_vehicle_and_new_customer_roll_back_on_file_error(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        customer, vehicle, record = self.make_archived_compact_policy()
        customer.delete()
        headers, row = self.compact_data()
        row[2] = "Bob"
        invalid = row.copy()
        invalid[5], invalid[6] = "MH12AB1235", "Unknown insurer"
        response = self.submit(headers, [row, invalid], template=template)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertEqual(Customer.all_objects.count(), 1)
        vehicle.refresh_from_db()
        self.assertEqual(vehicle.customer_id, customer.pk)
        self.assertIsNotNone(vehicle.deleted_at)
        record.refresh_from_db()
        self.assertIsNotNone(record.deleted_at)

    def test_archived_vehicle_reassignment_reuses_existing_uploaded_customer(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        customer, vehicle, record = self.make_archived_compact_policy()
        customer.delete()
        replacement = Customer.objects.create(name="Bob", phone="9876543212")
        headers, row = self.compact_data()
        row[2] = "Bob"
        preview = self.submit(headers, [row], template=template)
        imported = self.submit(headers, [row], template=template, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(InsuranceRecord.objects.get().customer_id, replacement.pk)
        self.assertEqual(Customer.objects.count(), 1)

    def test_deleted_customer_upload_is_restored_instead_of_skipped(self):
        customer = Customer.objects.create(name="Alice", phone="9876543210")
        customer.delete()
        headers, rows = ["Name", "Phone"], [["Alice", "9876543210"]]
        preview = self.submit(headers, rows)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["restored"], 1)
        self.assertEqual(preview.data["skipped"], 0)
        self.assertEqual(Customer.objects.count(), 0)
        imported = self.submit(headers, rows, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(Customer.objects.get().pk, customer.pk)

    def test_archived_vehicle_with_live_policy_cannot_change_owner(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        customer, vehicle, record = self.make_archived_compact_policy()
        Vehicle.all_objects.filter(pk=vehicle.pk).update(deleted_at=timezone.now())
        headers, row = self.compact_data()
        row[2] = "Bob"
        response = self.submit(headers, [row], template=template)
        self.assertEqual(response.status_code, 400, response.data)
        vehicle.refresh_from_db()
        self.assertEqual(vehicle.customer_id, customer.pk)
        self.assertIsNotNone(vehicle.deleted_at)

    def test_standard_archived_owner_reassignment_uses_uploaded_phone(self):
        customer, vehicle, record = self.make_archived_compact_policy("POL-1")
        customer.delete()
        headers, row = self.insurance_data()
        row[0] = "9876543212"
        headers.append("Customer Name")
        row.append("Bob")
        preview = self.submit(headers, [row], template=self.insurance)
        self.assertEqual(preview.status_code, 200, preview.data)
        imported = self.submit(headers, [row], template=self.insurance, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        restored = InsuranceRecord.objects.get()
        self.assertEqual(restored.pk, record.pk)
        self.assertEqual(restored.customer.name, "Bob")
        self.assertEqual(restored.customer.phone, "9876543212")
        self.assertEqual(restored.vehicle.customer_id, restored.customer_id)

    def test_standard_insurance_maps_net_pr_header(self):
        headers, row = self.insurance_data()
        headers.append("NET PR.")
        row.append("1000")
        preview = self.submit(headers, [row], template=self.insurance)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["warnings"], [])
        imported = self.submit(headers, [row], template=self.insurance, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(InsuranceRecord.objects.get().net_premium, 1000)

    def test_reimport_of_live_legacy_auto_policy_is_rejected(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        self.make_archived_compact_policy()
        headers, row = self.compact_data()
        response = self.submit(headers, [row], template=template)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("AUTO-OLD", str(response.data["errors"]))
        self.assertEqual(InsuranceRecord.objects.count(), 1)

    def test_archived_vehicle_can_be_reused_for_new_coverage_without_restoring_old_policy(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        customer, vehicle, old = self.make_archived_compact_policy()
        customer.delete()
        headers, row = self.compact_data()
        row[0], row[1] = "02.10.2026", "01.10.2027"
        preview = self.submit(headers, [row], template=template)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["created"], 1)
        imported = self.submit(headers, [row], template=template, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(InsuranceRecord.objects.get().vehicle_id, vehicle.pk)
        self.assertEqual(Customer.objects.count(), 1)
        old.refresh_from_db()
        self.assertIsNotNone(old.deleted_at)

    def test_archived_owner_name_formatting_differences_are_accepted(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        customer, vehicle, record = self.make_archived_compact_policy()
        customer.name = "SUSHMA A. DAWKHAR"
        customer.save()
        customer.delete()
        headers, row = self.compact_data()
        row[2] = "sushma  a\u00a0dawkhar"
        preview = self.submit(headers, [row], template=template)
        self.assertEqual(preview.status_code, 200, preview.data)
        imported = self.submit(headers, [row], template=template, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(InsuranceRecord.objects.get().customer_id, customer.pk)
        self.assertEqual(Customer.objects.count(), 1)

    def test_live_vehicle_uses_existing_owner_with_formatted_name_and_phone(self):
        customer = Customer.objects.create(name="SUSHMA A. DAWKHAR", phone="9876543210")
        vehicle = Vehicle.objects.create(customer=customer, vehicle_number="MH12AB1234", vehicle_type="Car")
        headers, row = self.insurance_data()
        headers.append("Customer Name")
        row.append("sushma a dawkhar")
        preview = self.submit(headers, [row], template=self.insurance)
        self.assertEqual(preview.status_code, 200, preview.data)
        imported = self.submit(headers, [row], template=self.insurance, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(InsuranceRecord.objects.get().customer_id, customer.pk)
        self.assertEqual(Customer.objects.count(), 1)

    def test_formatted_name_lookup_reuses_customer_for_a_second_vehicle(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        customer = Customer.objects.create(name="SUSHMA A. DAWKHAR", phone="9876543210")
        headers, row = self.compact_data()
        row[2] = "SUSHMA A DAWKHAR"
        preview = self.submit(headers, [row], template=template)
        self.assertEqual(preview.status_code, 200, preview.data)
        imported = self.submit(headers, [row], template=template, token=preview.data["preview_token"])
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(InsuranceRecord.objects.get().customer_id, customer.pk)
        self.assertEqual(Customer.objects.count(), 1)

    def test_real_owner_conflict_reports_both_names_without_changing_ownership(self):
        template = UploadTemplate.objects.get(name="Compact Insurance Records")
        customer, vehicle, record = self.make_archived_compact_policy()
        record.delete()  # The vehicle remains live; its owner cannot be replaced.
        headers, row = self.compact_data()
        row[2] = "Bob"
        response = self.submit(headers, [row], template=template)
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("Alice", str(response.data["errors"]))
        self.assertIn("Bob", str(response.data["errors"]))
        vehicle.refresh_from_db()
        self.assertEqual(vehicle.customer_id, customer.pk)
        self.assertIsNone(vehicle.deleted_at)
