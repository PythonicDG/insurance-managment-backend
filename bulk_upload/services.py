import hashlib
import io
import json
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.exceptions import ValidationError

from customers.models import Customer
from customers.serializers import CustomerSerializer
from insurance.models import InsuranceCompany, InsuranceRecord
from insurance.serializers import InsuranceRecordCreateUpdateSerializer
from vehicles.models import Vehicle
from .schema import FIELDS, normalize_header, validate_columns

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 100


def configuration(template):
    columns = list(template.columns.all())
    try:
        validate_columns(template.target, columns)
    except DjangoValidationError as exc:
        raise ValidationError({"error": "Template configuration is invalid. Contact your administrator.", "details": exc.messages})
    return [c for c in columns if c.is_active]


def fingerprint(template, columns):
    payload = [template.pk, template.target, template.is_active, [
        [c.pk, c.column_name, c.field_name, c.aliases, c.is_required, c.default_value, c.position]
        for c in columns
    ]]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False).encode()).hexdigest()


def convert(value, kind):
    if kind == "text":
        if isinstance(value, bool):
            raise ValueError("Enter text, not a Boolean value.")
        if isinstance(value, (int, float, Decimal)):
            number = Decimal(str(value))
            if not number.is_finite():
                raise ValueError("Enter a finite value.")
            value = format(number, "f")
            if "." in value:
                value = value.rstrip("0").rstrip(".")
        return str(value).strip()
    if kind == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
            try:
                return datetime.strptime(str(value).strip(), fmt).date().isoformat()
            except ValueError:
                continue
        raise ValueError("Use an Excel date, YYYY-MM-DD, DD/MM/YYYY, DD-MM-YYYY, or DD.MM.YYYY.")
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError("Enter a valid number without currency symbols or commas.")
    if not number.is_finite():
        raise ValueError("Enter a finite number.")
    return str(number)


def read_upload(upload):
    if not upload or not upload.name.lower().endswith(".xlsx"):
        raise ValidationError({"error": "Please upload an .xlsx Excel file."})
    if upload.size > MAX_FILE_BYTES:
        raise ValidationError({"error": "File exceeds the 10 MB limit."})
    raw = upload.read(MAX_FILE_BYTES + 1)
    if len(raw) > MAX_FILE_BYTES:
        raise ValidationError({"error": "File exceeds the 10 MB limit."})
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValidationError({"error": "Excel file expands beyond the supported size."})
    except BadZipFile:
        raise ValidationError({"error": "The file is not a valid .xlsx workbook."})
    return raw


def parse_rows(raw, template, columns):
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        sheet = workbook.worksheets[0]
        # Do not trust dimensions supplied by third-party Excel generators.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        header = next(iterator, ())
        if not header or len(header) > MAX_COLUMNS:
            raise ValidationError({"error": "Provide a header row with at most 100 columns on the first worksheet."})
        headers = [normalize_header(cell.value) for cell in header]
        nonempty = [name for name in headers if name]
        if len(nonempty) != len(set(nonempty)):
            raise ValidationError({"error": "The Excel file contains duplicate column headers."})
        positions, recognized = {}, set()
        for column in columns:
            names = {normalize_header(column.column_name), *[normalize_header(a) for a in column.aliases.splitlines() if a.strip()]}
            matches = [i for i, name in enumerate(headers) if name and name in names]
            if len(matches) > 1:
                raise ValidationError({"error": f"More than one header matches '{column.column_name}'. Keep only one."})
            if not matches and column.is_required and not column.default_value.strip():
                raise ValidationError({"error": f"Required column '{column.column_name}' is missing."})
            positions[column.pk] = matches[0] if matches else None
            recognized.update(matches)
        warnings = [f"Ignored column: {cell.value}" for i, cell in enumerate(header) if headers[i] and i not in recognized]
        rows = []
        for number, cells in enumerate(iterator, start=2):
            if number > MAX_ROWS + 1:
                raise ValidationError({"error": f"Use at most {MAX_ROWS} data rows per file, including blank rows."})
            if len(cells) > MAX_COLUMNS:
                raise ValidationError({"error": "Use at most 100 columns per file."})
            if all(cell.value in (None, "") for cell in cells):
                continue
            data, errors = {}, {}
            for column in columns:
                index = positions[column.pk]
                cell = cells[index] if index is not None and index < len(cells) else None
                value = cell.value if cell else None
                if cell and cell.data_type in ("f", "e"):
                    errors[column.column_name] = "Replace formulas or Excel errors with a plain value."
                    continue
                if value is None or (isinstance(value, str) and not value.strip()):
                    value = column.default_value.strip() or None
                if value is None:
                    if column.is_required:
                        errors[column.column_name] = "This value is required."
                    continue
                try:
                    data[column.field_name] = convert(value, FIELDS[template.target][column.field_name][1])
                except (ValueError, TypeError) as exc:
                    errors[column.column_name] = str(exc)
            rows.append((number, data, errors))
        if not rows:
            raise ValidationError({"error": "The first worksheet has no data rows."})
        return rows, warnings
    except ValidationError:
        raise
    except Exception:
        raise ValidationError({"error": "The workbook could not be read. Save it as a valid .xlsx file and try again."})
    finally:
        if workbook:
            workbook.close()


def prepare_insurance_data(data):
    """Adapt compact spreadsheets without relaxing the normal record API."""
    data = data.copy()
    net = data.pop("net_premium", None)
    if net is not None:
        net = Decimal(net)
        total = Decimal(data["total_premium"])
        if net < 0 or net > total or net.as_tuple().exponent < -2:
            raise ValidationError({"net_premium": "Net premium must be between zero and total premium, with at most two decimal places."})
        discount = total - net
        if "discount" in data and Decimal(data["discount"]) != discount:
            raise ValidationError({"net_premium": "Net premium must equal total premium minus discount."})
        data["discount"] = str(discount)
    if not data.get("policy_number"):
        # Stable across preview/import and repeated uploads, independent of row order.
        identity = [Vehicle.normalize_vehicle_number(data.get("vehicle_number", "")),
                    data.get("policy_start_date"), data.get("policy_expiry_date")]
        digest = hashlib.sha256(json.dumps(identity).encode()).hexdigest()[:32].upper()
        data["policy_number"] = f"BULK-{digest}"
    return data


def customer_name_key(name):
    """Treat spacing and punctuation in spreadsheet names as formatting."""
    return "".join(c for c in unicodedata.normalize("NFKC", name or "").casefold() if c.isalnum())


def prepare_placeholder_vehicle(data, notices):
    """NEW labels are temporary references, not globally unique registrations."""
    number = Vehicle.normalize_vehicle_number(data.get("vehicle_number", ""))
    if not re.fullmatch(r"NEW(?:\([^()]+\))?", number):
        return data
    owner = customer_name_key(data.get("customer_name", "")) or Customer.normalize_phone(data.get("customer_phone", ""))
    if not owner:
        raise ValidationError({"customer_name": "A customer name or mobile number is required for a NEW vehicle placeholder."})
    digest = hashlib.sha256(json.dumps([number, owner]).encode()).hexdigest()[:16].upper()
    data = data.copy()
    data["vehicle_number"] = f"{number[:30]}-{digest}"
    notices.append(f"Temporary vehicle '{number}' uses '{data['vehicle_number']}' for this customer. Replace it with the actual registration when available.")
    return data


def check_vehicle_owner(customer, data):
    name = data.get("customer_name", "").strip()
    phone = Customer.normalize_phone(data.get("customer_phone", ""))
    if not (customer_name_key(name) or phone) or (
        name and customer_name_key(customer.name) != customer_name_key(name)
    ) or (phone and customer.phone and phone != Customer.normalize_phone(customer.phone)):
        raise ValidationError({"vehicle_number": (
            f"Vehicle '{data.get('vehicle_number', '')}' owner is '{customer.name or '(name missing)'}', but the uploaded customer is "
            f"'{name or '(name missing)'}'. Customer details do not match; review this row."
        )})


def resolve_customer_without_phone(data, *, use_vehicle=True):
    name = data.get("customer_name", "").strip()
    if not customer_name_key(name) or len(name) > 255:
        raise ValidationError({"customer_name": "Enter a customer name of at most 255 characters when mobile number is missing."})
    vehicle_number = Vehicle.normalize_vehicle_number(data.get("vehicle_number", ""))
    vehicle = Vehicle.objects.select_related("customer").filter(vehicle_number__iexact=vehicle_number).first() if use_vehicle else None
    if vehicle:
        check_vehicle_owner(vehicle.customer, data)
        return vehicle.customer
    key = customer_name_key(name)
    matches = []
    for customer in Customer.objects.only("id", "name"):
        if customer_name_key(customer.name) == key:
            matches.append(customer)
            if len(matches) == 2:
                break
    if len(matches) > 1:
        raise ValidationError({"customer_name": "Multiple customers have this name. Supply a mobile number or review the existing vehicle owner."})
    if matches:
        return matches[0]
    return Customer.objects.create(name=name, phone="")


def check_archived_owner(customer, data):
    check_vehicle_owner(customer, data)


def revive_row(row):
    """Recover only this row, never its entire deletion batch or children."""
    if row.deleted_at is not None:
        type(row).all_objects.filter(pk=row.pk).update(deleted_at=None, deletion_batch=None)
        row.refresh_from_db()


def resolve_uploaded_customer(data):
    """Resolve a replacement owner without using the archived vehicle's owner."""
    phone = Customer.normalize_phone(data.get("customer_phone", ""))
    if not phone:
        return resolve_customer_without_phone(data, use_vehicle=False)
    values = {field: data[f"customer_{field}"] for field in ("name", "phone", "email", "address")
              if f"customer_{field}" in data}
    serializer = CustomerSerializer(data=values)
    serializer.is_valid(raise_exception=True)
    matches = Customer.objects.filter(phone=phone)
    name = data.get("customer_name", "")
    if name:
        matches = [c for c in matches if customer_name_key(c.name) == customer_name_key(name)]
    else:
        matches = list(matches[:2])
    if len(matches) > 1:
        raise ValidationError({"customer_phone": "Multiple customers match these details. Review the customer records."})
    return matches[0] if matches else serializer.save()


def recover_insurance_links(data, notices):
    vehicle = Vehicle.all_objects.select_for_update().select_related("customer").filter(
        vehicle_number__iexact=Vehicle.normalize_vehicle_number(data.get("vehicle_number", ""))
    ).first()
    if vehicle:
        customer = Customer.all_objects.select_for_update().get(pk=vehicle.customer_id)
        try:
            check_vehicle_owner(customer, data)
        except ValidationError:
            # A deleted vehicle reserves its number, but its archived owner must
            # not block a fresh upload. Never reassign a vehicle with live data.
            if vehicle.deleted_at is None or InsuranceRecord.objects.filter(vehicle=vehicle).exists():
                raise
            original_name = customer.name
            customer = resolve_uploaded_customer(data)
            vehicle.customer = customer
            vehicle.save(update_fields=["customer", "updated_at"])
            vehicle._bulk_owner_changed = True
            notices.append(f"Archived vehicle {vehicle.vehicle_number}: owner changed from '{original_name}' to '{customer.name}' using the uploaded customer.")
        revive_row(customer)
        revive_row(vehicle)
        data["customer_id"] = customer.pk
    return vehicle


def archived_policy(data, vehicle, supplied_policy):
    records = InsuranceRecord.all_objects.select_for_update().filter(deleted_at__isnull=False)
    if supplied_policy:
        record = records.filter(policy_number__iexact=data["policy_number"]).first()
        if record and (not vehicle or record.vehicle_id != vehicle.pk):
            raise ValidationError({"policy_number": "This archived policy number belongs to a different vehicle."})
        return record
    if not vehicle:
        return None
    # Also recognize references generated by earlier AUTO- import versions.
    matches = list(InsuranceRecord.all_objects.select_for_update().filter(
        vehicle=vehicle, insurance_company_id=data["insurance_company_id"],
        policy_start_date=data["policy_start_date"], policy_expiry_date=data["policy_expiry_date"],
    )[:2])
    if len(matches) > 1:
        raise ValidationError({"policy_number": "Multiple policies match this vehicle and coverage. Supply the policy number using the standard template."})
    if matches and matches[0].deleted_at is None:
        # Preserve duplicate validation even for old AUTO- references on expired policies.
        data["policy_number"] = matches[0].policy_number
        return None
    return matches[0] if matches else None


def import_row(target, data, notices=None):
    notices = notices if notices is not None else []
    if target == "customers":
        serializer = CustomerSerializer(data=data)
        serializer.is_valid(raise_exception=True)
        values = serializer.validated_data
        existing = Customer.objects.filter(phone=values["phone"])
        if values.get("name"):
            existing = existing.filter(name__iexact=values["name"])
        if existing.exists():
            return "skipped"
        archived = Customer.all_objects.filter(deleted_at__isnull=False, phone=values["phone"])
        if values.get("name"):
            archived = archived.filter(name__iexact=values["name"])
        matches = list(archived[:2])
        if len(matches) > 1:
            raise ValidationError({"phone": "Multiple deleted customers match these details. Review the archived customers."})
        if matches:
            serializer = CustomerSerializer(instance=matches[0], data=data)
            serializer.is_valid(raise_exception=True)
            serializer.save(deleted_at=None, deletion_batch=None)
            return "restored"
        serializer.save()
        return "created"
    supplied_policy = bool(data.get("policy_number"))
    data = prepare_placeholder_vehicle(data, notices)
    data = prepare_insurance_data(data)
    company_name = data.pop("insurance_company_name", "")
    companies = list(InsuranceCompany.objects.filter(name__iexact=company_name, is_active=True)[:2])
    if len(companies) != 1:
        raise ValidationError({"insurance_company_name": "Enter the exact name of one existing active insurance company."})
    data["insurance_company_id"] = companies[0].pk
    vehicle = recover_insurance_links(data, notices)
    archived = archived_policy(data, vehicle, supplied_policy)
    if archived:
        if not getattr(vehicle, "_bulk_owner_changed", False):
            customer = Customer.all_objects.select_for_update().get(pk=archived.customer_id)
            check_archived_owner(customer, data)
            if vehicle.customer_id != customer.pk:
                raise ValidationError({"vehicle_number": "The archived policy and vehicle have different customer owners."})
            revive_row(customer)
            data["customer_id"] = customer.pk
        data["policy_number"] = archived.policy_number
    if not data.get("customer_phone") and not data.get("customer_id"):
        data["customer_id"] = resolve_customer_without_phone(data).pk
    serializer = InsuranceRecordCreateUpdateSerializer(instance=archived, data=data)
    serializer.is_valid(raise_exception=True)
    if archived:
        # Do not bypass current/scheduled policy conflicts by using update().
        InsuranceRecord.activate_due_scheduled(vehicle_id=vehicle.pk)
        for old in InsuranceRecord.objects.filter(vehicle=vehicle, is_active=True, policy_expiry_date__lt=timezone.localdate()):
            old.is_active = False
            old.save(update_fields=["is_active"])
        if InsuranceRecord.objects.filter(vehicle=vehicle, is_active=True).exists() or InsuranceRecord.objects.filter(
            vehicle=vehicle, is_active=False, policy_start_date__gt=timezone.localdate(), renewed_policy__isnull=True,
        ).exists():
            raise ValidationError({"vehicle_number": "An active or scheduled policy already exists for this vehicle. Review or renew it before restoring this row."})
        start = serializer.validated_data["policy_start_date"]
        expiry = serializer.validated_data["policy_expiry_date"]
        record = serializer.save(deleted_at=None, deletion_batch=None,
                                 is_active=start <= timezone.localdate() <= expiry)
    else:
        record = serializer.save()
    if record.vehicle.customer_id != record.customer_id:
        raise ValidationError({"vehicle_number": "This vehicle belongs to a different customer. Review the customer details."})
    return "restored" if archived else "created"


def process_rows(target, rows, columns, preview):
    """Exercise the existing creation rules, rolling back previews and failed files."""
    errors, sample, warnings = [], [], []
    seen_vehicles = {}
    counts = {"created": 0, "skipped": 0, "restored": 0}
    names = {c.field_name: c.column_name for c in columns}
    with transaction.atomic():
        for number, data, cell_errors in rows:
            try:
                with transaction.atomic():
                    if cell_errors:
                        raise ValidationError(cell_errors)
                    notices = []
                    result = import_row(target, data, notices)
                    counts[result] += 1
                    warnings.extend(f"Excel row {number}: {notice}" for notice in notices)
                    if target == "insurance":
                        seen_vehicles.setdefault(Vehicle.normalize_vehicle_number(data.get("vehicle_number", "")), number)
                    if len(sample) < 10:
                        sample.append({"row": number, "status": result, "values": {names.get(k, k): v for k, v in data.items()}})
            except (ValidationError, DjangoValidationError, IntegrityError) as exc:
                if isinstance(exc, ValidationError):
                    detail = exc.detail
                elif isinstance(exc, DjangoValidationError):
                    detail = exc.message_dict if hasattr(exc, "message_dict") else exc.messages
                else:
                    detail = {"record": "A duplicate or conflicting record was detected. Review this row."}
                if isinstance(detail, dict) and "vehicle_number" in detail:
                    earlier = seen_vehicles.get(Vehicle.normalize_vehicle_number(data.get("vehicle_number", "")))
                    if earlier:
                        detail["vehicle_number"] = f"{detail['vehicle_number']} The same vehicle identifier also appears in Excel row {earlier}; check both rows."
                if isinstance(detail, dict):
                    detail = {names.get(k, k): v for k, v in detail.items()}
                errors.append({"row": number, "errors": detail})
        if preview or errors:
            transaction.set_rollback(True)
    return {"valid": not errors, "total_rows": len(rows), "created": counts["created"] if not errors or preview else 0,
            "restored": counts["restored"] if not errors or preview else 0,
            "skipped": counts["skipped"], "errors": errors, "sample": sample, "warnings": warnings}


def sample_workbook(template, columns):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Upload"
    sheet.append([c.column_name for c in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns, 1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        sheet.column_dimensions[cell.column_letter].width = max(20, min(45, len(column.column_name) + 5))
        # Preformat text fields to preserve leading zeros in phone/policy numbers.
        for row in range(2, 102):
            sheet.cell(row, index).number_format = "yyyy-mm-dd" if FIELDS[template.target][column.field_name][1] == "date" else "@" if FIELDS[template.target][column.field_name][1] == "text" else "0.00"
    help_sheet = workbook.create_sheet("Instructions")
    help_sheet.append(["Upload template", template.name])
    help_sheet.append(["Instructions", "Fill the first worksheet. Headers are on row 1. Maximum 2,000 data rows / 10 MB. No formulas."])
    help_sheet.append(["Dates", "Excel dates, YYYY-MM-DD, DD/MM/YYYY, DD-MM-YYYY, or DD.MM.YYYY (day first)."])
    if template.target == "insurance":
        help_sheet.append(["Deleted records", "Matching archived policies and customer/vehicle links are recovered. Other archived records and payments remain archived."])
        help_sheet.append(["Archived owner changes", "A deleted vehicle with no live policies can use the uploaded customer. Preview warnings show the old and new owners; review before importing."])
        help_sheet.append(["Missing details", "Without a mobile number, customer name is required. Missing policy numbers receive a stable BULK- reference."])
        help_sheet.append(["Net premium", "Net premium equals total premium minus discount. If supplied, the discount is calculated from these values."])
    help_sheet.append(["Duplicates", "Existing customers (same phone/name) are skipped. Duplicate policies are rejected."])
    help_sheet.append(["Insurance company", "Use the exact name of an existing active company from Settings."])
    help_sheet.append(["Column", "Required", "Type", "Default", "Accepted alternative headers"])
    for column in columns:
        help_sheet.append([column.column_name, "Yes" if column.is_required else "No", FIELDS[template.target][column.field_name][1], column.default_value, column.aliases.replace("\n", ", ")])
    # Force configurable strings to be text, even when they start with '='.
    for worksheet in workbook:
        for row in worksheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    help_sheet.column_dimensions["A"].width = 30
    help_sheet.column_dimensions["B"].width = 100
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
